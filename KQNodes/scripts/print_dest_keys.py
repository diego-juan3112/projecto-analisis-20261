from openpyxl import load_workbook

p='docs/DatosPruebas2026_1.xlsx'
wb=load_workbook(p)
sh='10A-Elementos'
if sh not in wb.sheetnames:
    sh=wb.sheetnames[0]
ws=wb[sh]
print('sheet', sh, 'rows', ws.max_row)

def normalize(s):
    if s is None: return ''
    s=str(s)
    # remove pipes, spaces, commas, empty symbol
    s=s.replace('|','').replace(' ', '').replace(',', '').replace('∅','')
    return ''.join(ch for ch in s if ch.isalpha()).upper()

for r in range(1, min(61, ws.max_row+1)):
    b=ws.cell(row=r, column=2).value
    c=ws.cell(row=r, column=3).value
    print(r, repr(b), '=>', normalize(b), ' | ', repr(c), '=>', normalize(c))
