from pathlib import Path
import re
from openpyxl import load_workbook

ABECEDARY = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

def bin_to_letters(bin_str):
    if bin_str is None:
        return ["", "", ""]
    s = str(bin_str).strip()
    letters = [ABECEDARY[i] for i, ch in enumerate(s) if ch == "1"]
    joined = "".join(letters)
    commas = ",".join(letters)
    commas_sp = ", ".join(letters)
    return [joined, commas, commas_sp]


def normalize_for_compare(s):
    if s is None:
        return ""
    s = str(s).strip()
    s = s.replace("|", "")
    s = s.replace("\u2217", "")
    s = s.replace("∅", "")
    s = s.replace(" ", "")
    s = s.replace(",", "")
    return s.upper()


BASE = Path.cwd()
res_path = BASE / "GeoMIP" / "results" / "resultados_Geometric.xlsx"
pruebas_path = BASE / "docs" / "DatosPruebas2026_1.xlsx"
print('res_path', res_path)
print('pruebas_path', pruebas_path)

wb_res = load_workbook(res_path, data_only=True)
# find sheet with Alcance/Mecanismo
ws_res = None
for name in wb_res.sheetnames:
    ws = wb_res[name]
    # find header row
    hdr = 1
    found = False
    for r in range(1, min(6, ws.max_row)+1):
        rowvals = [str((ws.cell(row=r, column=c).value or "")).upper() for c in range(1, min(20, ws.max_column)+1)]
        if any('ALCANCE' in v for v in rowvals) and any('MECANISMO' in v for v in rowvals):
            hdr = r
            found = True
            break
    if found:
        ws_res = ws
        header_row_res = hdr
        break
if ws_res is None:
    ws_res = wb_res[wb_res.sheetnames[0]]
    header_row_res = 1

# map columns
hdr_map_res = {}
for c in range(1, ws_res.max_column+1):
    v = ws_res.cell(row=header_row_res, column=c).value
    if v is None:
        continue
    hdr_map_res[str(v).strip().upper()] = c
col_alc_res = hdr_map_res.get('ALCANCE')
col_mec_res = hdr_map_res.get('MECANISMO')
print('col_alc_res', col_alc_res, 'col_mec_res', col_mec_res)

results = []
for r in range(header_row_res+1, header_row_res+1+50):
    if r>ws_res.max_row: break
    alc = ws_res.cell(row=r, column=col_alc_res).value if col_alc_res else None
    mec = ws_res.cell(row=r, column=col_mec_res).value if col_mec_res else None
    results.append((r, alc, mec))

print('\nSample resultados (first 20):')
for r,alc,mec in results[:20]:
    print(r, 'alc_raw=', repr(alc), 'mec_raw=', repr(mec))
    print('  bin->letters:', bin_to_letters(alc))
    print('  normalized bin variants:', [normalize_for_compare(x) for x in bin_to_letters(alc)])
    print('  normalized mec variants:', [normalize_for_compare(x) for x in bin_to_letters(mec)])

# load destination sheet
wb_pr = load_workbook(pruebas_path)
sheet_name = '10A-Elementos'
if sheet_name not in wb_pr.sheetnames:
    sheet_name = wb_pr.sheetnames[0]
ws_dest = wb_pr[sheet_name]
print('\nUsing destination sheet:', sheet_name)

dest_samples = []
dest_map = {}
for r in range(1, min(200, ws_dest.max_row)+1):
    b = ws_dest.cell(row=r, column=2).value
    c = ws_dest.cell(row=r, column=3).value
    key = (normalize_for_compare(b), normalize_for_compare(c))
    dest_samples.append((r, b, c, key))
    if key not in dest_map:
        dest_map[key] = r

print('\nSample dest rows (first 20):')
for row,b,c,key in dest_samples[:20]:
    print(row, 'B=', repr(b), 'C=', repr(c), 'key=', key)

# attempt match counts
matched = 0
notmatched = 0
misses = []
for r, alc, mec in results:
    alc_vars = bin_to_letters(alc)
    mec_vars = bin_to_letters(mec)
    found = False
    for a in alc_vars:
        for m in mec_vars:
            key = (normalize_for_compare(a), normalize_for_compare(m))
            if key in dest_map:
                found = True
                break
        if found:
            break
    if found:
        matched += 1
    else:
        notmatched += 1
        misses.append((r, alc, mec, [normalize_for_compare(x) for x in alc_vars], [normalize_for_compare(x) for x in mec_vars]))

print('\nResults sample matched:', matched, 'not matched:', notmatched)
print('\nFirst 10 misses:')
for m in misses[:10]:
    print(m)

# also show some dest_map keys present
print('\nDistinct dest keys sample (first 20):')
for i,k in enumerate(list(dest_map.keys())[:20]):
    print(i, k)

print('\nDone')
