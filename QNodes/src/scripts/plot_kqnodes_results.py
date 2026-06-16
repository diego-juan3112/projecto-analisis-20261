"""
Motor de Visualización Avanzado - Análisis Comparativo de Escalabilidad y IDCR
"""

import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from pathlib import Path
import warnings

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")

SCRIPT_DIR = Path(__file__).resolve().parent
SRC_DIR = SCRIPT_DIR.parent
EXCEL_PATH = SRC_DIR / ".samples" / "DatosPruebas2026_1JSME.xlsx"
OUTPUT_DIR = SRC_DIR / "results" / "plots"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

SHEETS = ["10A-Elementos", "15B-Elementos", "20A-Elementos", "22A-Elementos", "25A-Elementos "]

CONFIG_ESTRATEGIAS = {
    "Bipartición (Original)": {"loss_col": 5, "K": 2},  # Columna E
    "3-Partición (KQNodes)":  {"loss_col": 11, "K": 3}, # Columna K
    "4-Partición (KQNodes)":  {"loss_col": 17, "K": 4}, # Columna Q
    "5-Partición (KQNodes)":  {"loss_col": 23, "K": 5}  # Columna W
}

def cargar_datos_rendimiento():
    import openpyxl
    if not EXCEL_PATH.exists():
        return pd.DataFrame()
        
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    registros = []
    
    for sheet_name in SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        
        # Extraer número de nodos desde el nombre de la hoja
        n_nodos = int("".join(filter(str.isdigit, sheet_name)))
        
        max_row = ws.max_row
        while max_row > 7 and ws.cell(row=max_row, column=1).value is None:
            max_row -= 1
            
        for row in range(7, max_row + 1):
            for est_name, conf in CONFIG_ESTRATEGIAS.items():
                val_loss = ws.cell(row=row, column=conf["loss_col"]).value
                if val_loss is not None:
                    try:
                        registros.append({
                            "Nodos": n_nodos,
                            "Pestaña": sheet_name,
                            "Estrategia / Corte": est_name,
                            "K": conf["K"],
                            "Pérdida Integrada (IDCR)": float(val_loss)
                        })
                    except ValueError:
                        continue
    return pd.DataFrame(registros)

def main():
    print("📈 Extrayendo métricas consolidadas del repositorio Excel...")
    df = cargar_datos_rendimiento()
    
    if df.empty:
        print("❌ Error: No se encontraron datos procesados para graficar.")
        return

    sns.set_theme(style="whitegrid")
    
    # Gráfica 1: Comparativa de barras agrupadas
    plt.figure(figsize=(12, 6))
    df_agrupado = df.groupby(["Pestaña", "Estrategia / Corte"])["Pérdida Integrada (IDCR)"].mean().reset_index()
    
    ax = sns.barplot(
        data=df_agrupado,
        x="Pestaña",
        y="Pérdida Integrada (IDCR)",
        hue="Estrategia / Corte",
        palette="muted"
    )
    
    plt.title("Comparativa de Pérdida por Estrategia y Hoja (K-Particiones)", fontsize=13, fontweight='bold', pad=15)
    plt.xlabel("Muestra de Datos del Sistema (Pestañas)", fontweight='bold')
    plt.ylabel("Pérdida Promedio Evaluada", fontweight='bold')
    plt.legend(title="Configuración de Corte")
    
    plot1_path = OUTPUT_DIR / "comparativa_perdida_IDCR.png"
    plt.tight_layout()
    plt.savefig(plot1_path, dpi=150)
    plt.close()
    print(f"📊 Gráfica de barras exportada con éxito en: {plot1_path}")

    # Gráfica 2: Líneas de Tendencia de Escalabilidad
    plt.figure(figsize=(11, 6))
    sns.lineplot(
        data=df, 
        x="Nodos", 
        y="Pérdida Integrada (IDCR)", 
        hue="Estrategia / Corte", 
        marker="o", markersize=8,
        linewidth=2, palette="viridis"
    )
    
    plt.title("Tendencia del Índice de Disrupción Causacional Refinado (IDCR)", pad=15, fontweight='bold')
    plt.xlabel("Tamaño del Sistema (Nº de Nodos)", fontweight='bold')
    plt.ylabel("Gradiente de Pérdida Causal", fontweight='bold')
    plt.xticks([10, 15, 20, 22, 25])
    
    plot2_path = OUTPUT_DIR / "tendencia_IDCR.png"
    plt.tight_layout()
    plt.savefig(plot2_path, dpi=150)
    plt.close()
    print(f"📈 Gráfica de tendencia estructural guardada en: {plot2_path}")

if __name__ == "__main__":
    main()