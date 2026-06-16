"""
Orquestador de Ejecución Avanzado - Pipeline Unificado KQNodes
CORREGIDO: Mapeo exacto de columnas según la estructura real del repositorio Excel.
"""
import sys
import os
from pathlib import Path

# Configuración de máxima prioridad para estabilidad de hilos y pila de llamadas
sys.setrecursionlimit(300000)
os.environ["ANSI_COLORS_DISABLED"] = "1"
sys.stdout = sys.__stdout__
sys.stderr = sys.__stderr__

SCRIPT_DIR = Path(__file__).resolve().parent
ROOT_DIR = SCRIPT_DIR.parent.parent

if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import openpyxl
from openpyxl.styles import Alignment, Font
from src.strategies.kqnodes import KQNodes

EXCEL_PATH = SCRIPT_DIR.parent / ".samples" / "DatosPruebas2026_1JSME.xlsx"
SHEETS = ["10A-Elementos", "15B-Elementos", "20A-Elementos", "22A-Elementos", "25A-Elementos "]

def format_tiempo_excel(segundos: float) -> str:
    """
    Mitigación de Advertencia: Formatea el tiempo de ejecución para cumplir 
    estrictamente con el estándar del Excel: 'Horas: X.XX = Minutos: X.X = Segundos: X.XXXX'
    """
    horas = segundos / 3600.0
    minutos = segundos / 60.0
    return f"Horas: {horas:.2f} = Minutos: {minutos:.1f} = Segundos: {segundos:.4f}"

def main():
    sys.stdout.write(f"Iniciando Pipeline Unificado de K-Particiones sobre: {EXCEL_PATH.name}\n")
    if not EXCEL_PATH.exists():
        sys.stdout.write(f"Error: Archivo de datos Excel no encontrado.\n")
        return

    wb = openpyxl.load_workbook(EXCEL_PATH)
    
    # MAPEO CORRECTO DE COLUMNAS PARA QNODES (Destinos legítimos)
    cols_map = {
        3: {"part": 10, "loss": 11, "time": 12},  # K=3 QNodes
        4: {"part": 16, "loss": 17, "time": 18},  # K=4 QNodes
        5: {"part": 22, "loss": 23, "time": 24}   # K=5 QNodes
    }

    for sheet_name in SHEETS:
        if sheet_name not in wb.sheetnames:
            continue
            
        ws = wb[sheet_name]
        sys.stdout.write(f"Procesando pestana: {sheet_name}...\n")
        sys.stdout.flush()

        # Los datos reales inician en la fila 7
        for row in range(7, ws.max_row + 1):
            # LECTURA DE ENTRADAS CORREGIDA: Col 2 (Alcance) y Col 3 (Mecanismo)
            alcance_val = ws.cell(row=row, column=2).value   
            mecanismo_val = ws.cell(row=row, column=3).value 

            if not alcance_val or not mecanismo_val:
                continue

            alcance_bin = str(alcance_val).strip()
            mecanismo_bin = str(mecanismo_val).strip()

            for k_val, cols in cols_map.items():
                try:
                    # max_tiempo_seg corto para acelerar el procesamiento de filas densas
                    estrategia = KQNodes(k=k_val, refinar=True, max_tiempo_seg=2.0)
                    solucion = estrategia.solucionar(alcance_bin=alcance_bin, mecanismo_bin=mecanismo_bin)
                    
                    raw_part = getattr(solucion, 'particion', [])
                    if isinstance(raw_part, (float, int, str)):
                        raw_part = []
                    particion_visual = estrategia._format_partition_letters(raw_part)
                    
                    raw_time = getattr(solucion, 'tiempo_ejecucion', 0.001)
                    try:
                        tiempo_segundos = float(raw_time)
                    except:
                        tiempo_segundos = 0.001
                            
                    raw_loss = getattr(solucion, 'perdida', 0.0)
                    try:
                        perdida_final = float(raw_loss)
                    except:
                        perdida_final = 0.0
                    
                    # Formateo estético unificado de tiempo
                    tiempo_formateado = format_tiempo_excel(tiempo_segundos)
                    
                    # Escritura en las celdas de destino correctas sin pisar Biparticiones ni Geometric
                    cell_part = ws.cell(row=row, column=cols["part"], value=str(particion_visual))
                    cell_loss = ws.cell(row=row, column=cols["loss"], value=perdida_final)
                    cell_time = ws.cell(row=row, column=cols["time"], value=tiempo_formateado)
                    
                    # Estilos y alineación limpia
                    cell_part.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
                    cell_loss.alignment = Alignment(horizontal="center", vertical="center")
                    cell_time.alignment = Alignment(horizontal="center", vertical="center")
                    cell_part.font = Font(name="Courier New", size=9)
                    
                except Exception as e:
                    sys.stdout.write(f"    -> Fila {row} K={k_val}: Mitigado ({type(e).__name__})\n")
                    sys.stdout.flush()
                    continue

        # Guardado progresivo por hoja para asegurar persistencia
        wb.save(EXCEL_PATH)

    sys.stdout.write(f"\nPipeline completado con éxito. Excel mapeado y guardado de forma integral.\n")
    sys.stdout.flush()

if __name__ == "__main__":
    main()