"""
Rellena columnas de resultados GeoMIP (k=2 Geometric) en DatosPruebas2026_1.xlsx.

Uso:
    uv run python scripts/fill_excel_geomip.py

Configura las constantes de la sección CONFIGURACIÓN antes de ejecutar.
Los casos se ejecutan en paralelo (MAX_WORKERS hilos, cada uno gestiona su propio
subproceso con timeout). Si un caso supera MAX_TIME_SEG se termina el proceso y se
escribe "TIMEOUT" en Excel.
"""

import concurrent.futures
import multiprocessing
import re
import sys
import traceback
from pathlib import Path

import openpyxl

PROJECT_ROOT  = Path(__file__).resolve().parent.parent
GEOMIP_METHOD = PROJECT_ROOT / "GeoMIP" / "src" / "Method2_Dynamic_Programming_Reformulation"

# Inyectar site-packages del venv de GeoMIP para que numpy/scipy estén disponibles
# cuando se ejecuta desde el venv raíz del proyecto (que no los tiene).
_geomip_sp = GEOMIP_METHOD / ".venv" / "Lib" / "site-packages"
if _geomip_sp.exists():
    sys.path.insert(0, str(_geomip_sp))

sys.path.insert(0, str(GEOMIP_METHOD))   # necesario: GeoMIP usa "from src.*"

# ========================= CONFIGURACIÓN =========================
EXCEL_PATH   = PROJECT_ROOT / "docs" / "DatosPruebas2026_1.xlsx"
SHEET_NAME   = "22A-Elementos"
MAX_TIME_SEG = 7200.0   # segundos máximos por caso; si se supera se escribe "TIMEOUT"

# Número de hilos paralelos. Cada hilo gestiona su propio subproceso con timeout,
# por lo que se pueden correr MAX_WORKERS casos al mismo tiempo sin riesgo de
# bloquear el pool ante timeouts. En AMD 3020e (2 núcleos físicos) usar 2.
# OJO: cada worker carga su propia copia de la TPM en RAM. Para N grande la TPM
# no cabe dos veces en memoria, así que MAX_WORKERS se limita automáticamente
# según N (ver _workers_efectivos). Este valor es solo el tope deseado.
MAX_WORKERS = 2

# A partir de este N se fuerza 1 solo worker para evitar MemoryError al cargar
# dos copias de la TPM a la vez (N=22 → TPM float32 ~0.37 GB + n-cubos ~0.35 GB).
N_FORZAR_UN_WORKER = 22

# Columna destino en Excel (1-based).
# Estructura por grupo de 3 cols: Partición | Pérdida | Tiempo
#   k=2 QNodes    → cols  4,  5,  6  (D, E, F)
#   k=2 Geometric → cols  7,  8,  9  (G, H, I)  ← activo
#   k=3 QNodes    → cols 10, 11, 12  (J, K, L)
#   k=3 Geometric → cols 13, 14, 15  (M, N, O)
#   k=4 QNodes    → cols 16, 17, 18  (P, Q, R)
#   k=4 Geometric → cols 19, 20, 21  (S, T, U)
#   k=5 QNodes    → cols 22, 23, 24  (V, W, X)
#   k=5 Geometric → cols 25, 26, 27  (Y, Z, AA)
COL_INICIO = 7   # primera columna del grupo (Partición)

HEADER_ROW      = 5      # fila de encabezados de columnas en el Excel
SKIP_SI_RELLENO = True   # True = saltar filas ya rellenas (recomendado para re-ejecuciones)
# =================================================================


# ---------------------------------------------------------------------------
# Worker que corre en subproceso separado (obligatorio a nivel de módulo en Windows)
# ---------------------------------------------------------------------------
def _worker_geomip(
    queue: multiprocessing.Queue,
    geomip_method: str,
    geomip_sp: str,
    estado_inicial: str,
    pagina: str,
    condiciones: str,
    alcance_bin: str,
    mecanismo_bin: str,
) -> None:
    """Se ejecuta en un Process independiente. Envía (particion, perdida, tiempo) por queue."""
    try:
        sys.path.insert(0, geomip_sp)
        sys.path.insert(0, geomip_method)

        import numpy as np                                             # noqa: PLC0415
        import pandas as pd                                             # noqa: PLC0415

        from src.models.base.application import aplicacion              # noqa: PLC0415
        from src.controllers.manager import Manager                     # noqa: PLC0415
        from src.controllers.strategies.geometric import GeometricSIA  # noqa: PLC0415

        aplicacion.pagina_sample_network = pagina
        gestor    = Manager(estado_inicial)
        geo       = GeometricSIA(gestor)

        # Carga eficiente de la TPM. np.genfromtxt (sia_cargar_tpm) construye
        # listas de Python y su pico de RAM es ~5-10x el tamaño del archivo, lo
        # que provoca MemoryError con TPMs grandes (N>=22 → >1 GB). pandas usa el
        # parser en C y dtype=float32 directo: pico bajo y arreglo final pequeño.
        # float32 es idéntico a lo que System ya hace internamente (system.py:41).
        tpm = pd.read_csv(
            geo.sia_gestor.tpm_filename,
            header=None,
            dtype=np.float32,
        ).to_numpy()
        # Guard: si hay una coma final en el CSV, pandas añade una columna NaN.
        if tpm.shape[1] == len(estado_inicial) + 1 and np.isnan(tpm[:, -1]).all():
            tpm = np.ascontiguousarray(tpm[:, :-1])

        resultado = geo.aplicar_estrategia(condiciones, alcance_bin, mecanismo_bin, tpm)
        queue.put(("ok", str(resultado.particion), float(resultado.perdida), float(resultado.tiempo_ejecucion)))
    except Exception as exc:  # noqa: BLE001
        queue.put(("err", str(exc), traceback.format_exc()))


# ---------------------------------------------------------------------------
# Wrapper ejecutado por cada hilo del ThreadPoolExecutor.
# Crea su propio Process con timeout y lo termina si se excede.
# ---------------------------------------------------------------------------
def _run_geomip_con_timeout(
    geomip_method: str,
    geomip_sp: str,
    estado_inicial: str,
    pagina: str,
    condiciones: str,
    alcance_bin: str,
    mecanismo_bin: str,
) -> tuple:
    """Lanza GeoMIP en subproceso con timeout. Retorna la tupla del queue o una marca especial."""
    queue: multiprocessing.Queue = multiprocessing.Queue()
    proc = multiprocessing.Process(
        target=_worker_geomip,
        args=(queue, geomip_method, geomip_sp, estado_inicial, pagina,
              condiciones, alcance_bin, mecanismo_bin),
        daemon=True,
    )
    proc.start()
    proc.join(timeout=MAX_TIME_SEG)
    if proc.is_alive():
        proc.terminate()
        proc.join()
        return ("timeout",)
    if queue.empty():
        return ("crash",)
    return queue.get()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def leer_metadatos_hoja(ws) -> tuple[str, str]:
    estado_inicial = str(ws.cell(1, 2).value).strip()
    sistema        = str(ws.cell(2, 2).value).strip()
    return estado_inicial, sistema


def leer_casos_prueba(ws) -> list[dict]:
    casos = []
    for row_idx in range(HEADER_ROW + 1, ws.max_row + 1):
        alc = ws.cell(row_idx, 2).value
        mec = ws.cell(row_idx, 3).value
        if alc is None or mec is None:
            continue
        alc_s = str(alc).strip()
        mec_s = str(mec).strip()
        if not alc_s or not mec_s:
            continue
        casos.append({"row_idx": row_idx, "alcance_letras": alc_s, "mecanismo_letras": mec_s})
    return casos


def letras_a_binario(letras: str, sistema: str) -> str:
    return "".join("1" if c in letras else "0" for c in sistema)


def formatear_tiempo(tiempo_s: float) -> str:
    return (
        f"Horas: {tiempo_s / 3600:.2f} = "
        f"Minutos: {tiempo_s / 60:.1f} = "
        f"Segundos: {tiempo_s:.4f}"
    )


def escribir_resultado(
    excel_path: Path,
    sheet_name: str,
    row_idx: int,
    col_inicio: int,
    particion: str,
    perdida,
    tiempo: str,
) -> None:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name]
    ws.cell(row_idx, col_inicio    ).value = particion
    ws.cell(row_idx, col_inicio + 1).value = str(round(float(perdida), 4)) if isinstance(perdida, float) else perdida
    ws.cell(row_idx, col_inicio + 2).value = tiempo
    wb.save(excel_path)
    wb.close()


def workers_efectivos(n: int) -> int:
    """Limita los workers según N para no agotar la RAM cargando varias TPMs."""
    if n >= N_FORZAR_UN_WORKER:
        return 1
    return MAX_WORKERS


def extraer_pagina(sheet_name: str) -> str:
    m = re.match(r"\d+([A-Z])", sheet_name)
    if not m:
        raise ValueError(f"No se pudo extraer la página del nombre de hoja: {sheet_name!r}")
    return m.group(1)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    print(f"Excel      : {EXCEL_PATH}")
    print(f"Hoja       : {SHEET_NAME}")
    print(f"k          : 2  (bipartición GeoMIP Geometric)")
    print(f"Timeout    : {MAX_TIME_SEG}s por caso")
    print(f"Cols       : {COL_INICIO}–{COL_INICIO + 2} (Partición, Pérdida, Tiempo)")
    print()

    # 1 — Leer metadatos y casos del Excel
    wb_r = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_r = wb_r[SHEET_NAME]
    estado_inicial, sistema = leer_metadatos_hoja(ws_r)
    condiciones = "1" * len(estado_inicial)
    casos       = leer_casos_prueba(ws_r)
    # Una fila cuenta como "rellena" (y por tanto se omite) SOLO si su celda de
    # Partición tiene un valor real y NO es una marca de TIMEOUT. Las filas que
    # quedaron en timeout se tratan como faltantes y se vuelven a procesar.
    def _es_timeout(v) -> bool:
        return v is not None and str(v).strip().upper() == "TIMEOUT"

    filas_rellenas: set[int] = set()
    filas_timeout:  set[int] = set()
    if SKIP_SI_RELLENO:
        for row_idx in range(HEADER_ROW + 1, ws_r.max_row + 1):
            v = ws_r.cell(row_idx, COL_INICIO).value
            if _es_timeout(v):
                filas_timeout.add(row_idx)      # faltante: se reintenta
            elif v not in (None, ""):
                filas_rellenas.add(row_idx)     # resultado válido: se omite
    wb_r.close()

    pagina = extraer_pagina(SHEET_NAME)
    n_vars = len(estado_inicial)
    wef    = workers_efectivos(n_vars)

    print(f"Estado inicial : {estado_inicial}  (N={n_vars})")
    print(f"Sistema        : {sistema}")
    print(f"Condiciones    : {condiciones}  (todo 1s)")
    nota_w = "" if wef == MAX_WORKERS else f"  (limitado desde {MAX_WORKERS} por N grande / RAM)"
    print(f"Workers        : {wef}  (hilos paralelos, cada uno con su subproceso){nota_w}")
    print(f"Casos de prueba: {len(casos)}  (ya rellenos: {len(filas_rellenas)}, "
          f"TIMEOUT a reintentar: {len(filas_timeout)})")
    print(f"TPM            : N{len(estado_inicial)}{pagina}.csv")
    print()
    print("-" * 70)

    timeouts = 0
    errores  = 0

    # 2 — Preparar trabajos pendientes
    trabajos: list[dict] = []
    for i, caso in enumerate(casos):
        if caso["row_idx"] in filas_rellenas:
            print(f"[{i + 1}/{len(casos)}] fila={caso['row_idx']} → ya relleno, omitiendo")
            continue
        trabajos.append({
            "idx_global": i,
            "caso": caso,
            "alcance_bin":   letras_a_binario(caso["alcance_letras"],   sistema),
            "mecanismo_bin": letras_a_binario(caso["mecanismo_letras"], sistema),
        })

    if not trabajos:
        print("Nada que procesar — todos los casos ya estaban rellenos.")
        return

    print(f"\nProcesando {len(trabajos)} caso(s) con {wef} worker(s)…\n")

    # 3 — Ejecutar en paralelo con hilos; cada hilo gestiona su subproceso con timeout.
    #     La escritura al Excel se mantiene serial en el proceso principal.
    with concurrent.futures.ThreadPoolExecutor(max_workers=wef) as executor:
        future_to_trabajo = {
            executor.submit(
                _run_geomip_con_timeout,
                str(GEOMIP_METHOD), str(_geomip_sp),
                estado_inicial, pagina, condiciones,
                t["alcance_bin"], t["mecanismo_bin"],
            ): t
            for t in trabajos
        }

        completados = 0
        for future in concurrent.futures.as_completed(future_to_trabajo):
            t    = future_to_trabajo[future]
            caso = t["caso"]
            completados += 1

            print(
                f"[{completados}/{len(trabajos)}] fila={caso['row_idx']}  "
                f"Alcance={caso['alcance_letras']}  "
                f"Mecanismo={caso['mecanismo_letras']}"
            )

            try:
                msg = future.result()
            except Exception as e:
                errores += 1
                print(f"  ✗ Excepción en hilo: {type(e).__name__}: {e}\n")
                continue

            if msg[0] == "timeout":
                timeouts += 1
                print(f"  ⚠ TIMEOUT ({MAX_TIME_SEG}s) — se escribe marca en Excel\n")
                escribir_resultado(
                    EXCEL_PATH, SHEET_NAME, caso["row_idx"],
                    COL_INICIO, "TIMEOUT", "TIMEOUT", f">{MAX_TIME_SEG}s",
                )

            elif msg[0] == "crash":
                errores += 1
                print("  ✗ El proceso terminó sin resultado (crash silencioso)\n")

            elif msg[0] == "err":
                errores += 1
                _, exc_str, tb_str = msg
                sep = "=" * 70
                print(f"\n{sep}")
                print(f"ERROR en fila {caso['row_idx']}")
                print(f"  Alcance  : {caso['alcance_letras']} → {t['alcance_bin']}")
                print(f"  Mecanismo: {caso['mecanismo_letras']} → {t['mecanismo_bin']}")
                print(f"  Excepción: {exc_str}")
                print(tb_str)
                print(sep)
                print("Caso omitido — continuando con los demás.\n")

            else:
                _, particion_str, perdida, tiempo_ejec = msg
                tiempo_str = formatear_tiempo(tiempo_ejec)

                print(f"  Partición : {particion_str}")
                print(f"  Pérdida   : {perdida:.4f}")
                print(f"  {tiempo_str}")

                escribir_resultado(
                    EXCEL_PATH, SHEET_NAME, caso["row_idx"],
                    COL_INICIO, particion_str, perdida, tiempo_str,
                )
                print(f"  ✓ Escrito en Excel\n")

    print("=" * 70)
    print(f"Completado. Hoja='{SHEET_NAME}'  Procesados={len(trabajos)}  Timeouts={timeouts}  Errores={errores}")


if __name__ == "__main__":
    multiprocessing.freeze_support()   # necesario en Windows con PyInstaller / uv
    main()
