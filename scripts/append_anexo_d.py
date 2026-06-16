"""
Anexa el "Anexo D — Cobertura experimental y limitaciones de hardware" al final
de los manuales Word de KQNodes, SIN regenerarlos (preserva imágenes, el enlace
del video y todo el contenido existente que el usuario insertó a mano).

Idempotente: si el documento ya contiene un encabezado "Anexo D", lo omite.

Uso:
    uv run python scripts/append_anexo_d.py
"""

from pathlib import Path

import docx
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
TECNICO = ROOT / "KQNodes" / "docs" / "Manual_Tecnico_KQNodes.docx"
USUARIO = ROOT / "KQNodes" / "docs" / "Manual_Usuario_KQNodes.docx"

H2 = "Heading 2"
H3 = "Heading 3"
NOTE = "Block Text"
TABLE = "Table"


def ya_tiene_anexo_d(doc) -> bool:
    return any(p.text.strip().startswith("Anexo D") for p in doc.paragraphs)


def add_table(doc, headers, filas):
    t = doc.add_table(rows=1, cols=len(headers))
    try:
        t.style = TABLE
    except KeyError:
        pass
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = h
        for r in cell.paragraphs[0].runs:
            r.font.bold = True
    for fila in filas:
        cells = t.add_row().cells
        for j, val in enumerate(fila):
            cells[j].text = str(val)
    return t


def p(doc, texto, style=None):
    return doc.add_paragraph(texto, style=style)


# ---------------------------------------------------------------------------
def construir_tecnico(doc):
    doc.add_paragraph("Anexo D — Cobertura experimental y limitaciones de hardware", style=H2)
    p(doc,
      "El conjunto de pruebas de docs/DatosPruebas2026_1.xlsx incluye redes de 10, 15, "
      "20, 22 y 25 elementos. No todos los casos pudieron ejecutarse en el equipo de "
      "desarrollo por límites de memoria y tiempo. Esta sección documenta la causa, las "
      "mitigaciones aplicadas y la cobertura realmente alcanzada.")

    doc.add_paragraph("D.1 Causa raíz — costo exponencial en N", style=H3)
    p(doc,
      "El problema es intrínsecamente exponencial: un sistema de N nodos tiene 2^N estados, "
      "por lo que la TPM tiene 2^N filas × N columnas. Cada elemento adicional duplica (como "
      "mínimo) la memoria y el tiempo necesarios. Esto se refleja en el tamaño de los "
      "archivos de muestra (CSV de la TPM):")
    add_table(doc,
              ["N", "Estados (2^N)", "Tamaño TPM (CSV)"],
              [["10", "1 024", "~21 KB"],
               ["15", "32 768", "~5 MB"],
               ["20", "1 048 576", "~273 MB"],
               ["22", "4 194 304", "~1.2 GB"],
               ["25", "33 554 432", "~10.9 GB"]])

    doc.add_paragraph("D.2 Límite de memoria (RAM)", style=H3)
    p(doc,
      "Equipo de referencia: AMD 3020e (2 núcleos físicos a 1.2 GHz), RAM física limitada "
      "(~6 GB, con ~1.7 GB libres durante la ejecución).")
    p(doc,
      "• Cargar la TPM de N=22 con el método original (np.genfromtxt en sia_cargar_tpm) "
      "provocaba MemoryError: ese parser construye listas de Python y su pico de RAM es "
      "~5–10× el tamaño del archivo (>6 GB para 1.2 GB de CSV).")
    p(doc,
      "• Mitigación aplicada (scripts/fill_excel_geomip.py): se reemplazó la carga por "
      "pandas.read_csv(..., dtype=np.float32) (parser en C, pico bajo). float32 es "
      "numéricamente idéntico a lo que System ya hace internamente. Además se limita a 1 "
      "worker para N≥22, evitando dos copias simultáneas de la TPM en memoria.")
    p(doc,
      "• Con esto N=22 deja de fallar por memoria, pero N=25 sigue siendo inviable en este "
      "hardware: solo el arreglo final en float32 ocupa 2^25 × 25 × 4 bytes ≈ 3.3 GB, que ya "
      "excede la RAM libre, sin contar los n-cubos derivados ni el sistema operativo.")

    doc.add_paragraph("D.3 Límite de tiempo", style=H3)
    p(doc,
      "Cada caso se ejecuta con un timeout de 1 hora (MAX_TIME_SEG = 3600 s). Si se supera, "
      "el script escribe TIMEOUT en la celda y continúa con el siguiente caso (las celdas en "
      "TIMEOUT se reintentan en corridas posteriores). Para N grande, un único caso puede "
      "acercarse o superar ese límite, de modo que completar las ~50 pruebas de una hoja "
      "tomaría días en este equipo.")

    doc.add_paragraph("D.4 Cobertura realmente alcanzada", style=H3)
    add_table(doc,
              ["Hoja", "N", "k=2 QNodes", "k=2 Geometric", "k=3 / k=4 / k=5", "Estado"],
              [["10A", "10", "49/49", "49/49", "49 / 49 / 49", "Completo"],
               ["15B", "15", "50/50", "50/50", "50 / 50 / 50", "Completo"],
               ["20A", "20", "50/50", "50/50", "50 / 49 (1 timeout) / 50", "Prácticamente completo"],
               ["22A", "22", "32/50", "50/50", "0 / 0 / 0", "Parcial"],
               ["25A", "25", "0/50", "0/50", "0 / 0 / 0", "No ejecutado"]])

    doc.add_paragraph("D.5 Reproducibilidad y cómo completar lo pendiente", style=H3)
    p(doc,
      "• Los resultados de 10, 15 y 20 elementos están completos y se analizan en la hoja "
      "«Análisis 10-15-20» del Excel (tablas de φ y tiempos, comparación QNodes vs Geometric "
      "y crecimiento por k).")
    p(doc,
      "• Para completar 22A y 25A se requiere un equipo con más RAM (≥16–32 GB) y, "
      "preferiblemente, más núcleos; alternativamente subir el timeout y procesar por lotes. "
      "El script fill_excel_geomip.py reintenta automáticamente las celdas marcadas TIMEOUT, "
      "por lo que la corrida puede reanudarse sin perder el progreso ya escrito.")


def construir_usuario(doc):
    doc.add_paragraph("Anexo D — Por qué no se ejecutaron todos los casos del Excel", style=H2)
    p(doc,
      "Las pruebas de docs/DatosPruebas2026_1.xlsx cubren redes de 10, 15, 20, 22 y 25 "
      "elementos. No fue posible ejecutar todos los casos en el equipo de pruebas; aquí se "
      "explica por qué.")

    doc.add_paragraph("D.1 El costo crece de forma exponencial", style=H3)
    p(doc,
      "El esfuerzo para calcular φ crece exponencialmente con el número de elementos N: una "
      "red de N nodos tiene 2^N estados posibles, así que cada elemento extra duplica (o más) "
      "la memoria y el tiempo necesarios. Esto se ve directamente en el tamaño de los "
      "archivos de datos (la TPM):")
    add_table(doc,
              ["Red (N)", "Tamaño del archivo de datos"],
              [["10 elementos", "~21 KB"],
               ["15 elementos", "~5 MB"],
               ["20 elementos", "~273 MB"],
               ["22 elementos", "~1.2 GB"],
               ["25 elementos", "~10.9 GB"]])

    doc.add_paragraph("D.2 Las limitaciones del equipo", style=H3)
    p(doc,
      "El computador usado para las pruebas (AMD 3020e, 2 núcleos, RAM limitada) no alcanza "
      "para los casos más grandes:")
    p(doc,
      "• 10, 15 y 20 elementos: se completaron (con un único caso en timeout para k=4 en la "
      "red de 20).")
    p(doc,
      "• 22 elementos: se completó parcialmente — la bipartición Geometric (50/50) y parte "
      "de la bipartición QNodes (32/50); las particiones de k=3 a k=5 quedaron pendientes.")
    p(doc,
      "• 25 elementos: no fue posible ejecutarlo; el archivo de datos por sí solo (10.9 GB) "
      "ya no cabe en la memoria del equipo.")

    doc.add_paragraph("D.3 Qué se hizo y qué haría falta", style=H3)
    p(doc,
      "Para aprovechar el equipo al máximo se optimizó la carga de datos (formato más "
      "liviano, float32) y se limitó la ejecución a un proceso a la vez en las redes grandes "
      "para no agotar la memoria. Aun así, completar 22 y 25 elementos requiere un computador "
      "con más RAM (16 GB o más) y, de preferencia, más núcleos.")
    p(doc,
      "Los resultados de 10, 15 y 20 elementos —que sí están completos— se encuentran "
      "analizados en la hoja «Análisis 10-15-20» del mismo Excel, con sus tablas, gráficos e "
      "interpretación.")


def procesar(path: Path, constructor) -> None:
    doc = docx.Document(path)
    if ya_tiene_anexo_d(doc):
        print(f"[OMITIDO] {path.name} ya contiene 'Anexo D'.")
        return
    constructor(doc)
    doc.save(path)
    print(f"[OK] Anexo D agregado a {path.name}")


if __name__ == "__main__":
    procesar(TECNICO, construir_tecnico)
    procesar(USUARIO, construir_usuario)
