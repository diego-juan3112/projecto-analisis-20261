"""
Rellena columnas de resultados QNodes (k=2) en DatosPruebas2026_1.xlsx.

Uso:
    uv run python scripts/fill_excel_qnodes.py

Configura las constantes de la sección CONFIGURACIÓN antes de ejecutar.
Los casos se ejecutan en paralelo (MAX_WORKERS procesos simultáneos).
"""

import concurrent.futures
import re
import sys
import traceback
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parent.parent
QNODES_DIR   = PROJECT_ROOT / "QNodes"
QNODES_SRC   = QNODES_DIR / "src"

# Inyectar site-packages del venv de QNodes para que numpy/pyphi estén disponibles
# cuando se ejecuta desde el venv raíz del proyecto (que no los tiene).
_qnodes_sp = QNODES_DIR / ".venv" / "Lib" / "site-packages"
if _qnodes_sp.exists():
    sys.path.insert(0, str(_qnodes_sp))

sys.path.insert(0, str(QNODES_DIR))   # necesario: QNodes usa "from src.*"

# ========================= CONFIGURACIÓN =========================
EXCEL_PATH   = PROJECT_ROOT / "docs" / "DatosPruebas2026_1.xlsx"
SHEET_NAME   = "22A-Elementos"
MAX_TIME_SEG = 1800.0   # referencia informativa

# Número de procesos paralelos. Cada worker carga su PROPIA copia completa de
# la TPM en RAM. Para N=22 (~738 MB por copia) con solo 8 GB de RAM, usar 1
# worker evita el MemoryError de tener dos TPM cargadas a la vez. Para redes
# pequeñas (N≤15) se puede subir a 2.
MAX_WORKERS = 1

# Columna destino en Excel (1-based).
# Estructura por grupo de 3 cols: Partición | Pérdida | Tiempo
#   k=2 QNodes    → cols  4,  5,  6  (D, E, F)   ← activo
#   k=2 Geometric → cols  7,  8,  9  (G, H, I)
#   k=3 QNodes    → cols 10, 11, 12  (J, K, L)
#   k=3 Geometric → cols 13, 14, 15  (M, N, O)
#   k=4 QNodes    → cols 16, 17, 18  (P, Q, R)
#   k=4 Geometric → cols 19, 20, 21  (S, T, U)
#   k=5 QNodes    → cols 22, 23, 24  (V, W, X)
#   k=5 Geometric → cols 25, 26, 27  (Y, Z, AA)
COL_INICIO = 4   # primera columna del grupo (Partición)

HEADER_ROW      = 5      # fila de encabezados de columnas en el Excel
SKIP_SI_RELLENO = True   # True = saltar filas ya rellenas (recomendado con ejecución paralela)
# =================================================================


# ---------------------------------------------------------------------------
# Worker a nivel de módulo — obligatorio para multiprocessing en Windows
# ---------------------------------------------------------------------------
def _ejecutar_caso_qnodes(
    qnodes_dir: str,
    qnodes_sp: str,
    estado_inicial: str,
    pagina: str,
    condiciones: str,
    alcance_bin: str,
    mecanismo_bin: str,
) -> tuple[str, float, float]:
    """Corre QNodes en un proceso hijo independiente. Retorna (particion, perdida, tiempo)."""
    sys.path.insert(0, qnodes_sp)
    sys.path.insert(0, qnodes_dir)

    from src.models.base.application import aplicacion       # noqa: PLC0415
    from src.controllers.manager import Manager              # noqa: PLC0415
    from src.strategies.q_nodes import QNodes                # noqa: PLC0415
    from pathlib import Path as _Path                        # noqa: PLC0415

    aplicacion.set_pagina_red_muestra(pagina)
    sample_path = _Path(qnodes_dir) / "src" / ".samples"
    gestor   = Manager(estado_inicial, ruta_base=sample_path)
    tpm      = gestor.cargar_red()
    q        = QNodes(tpm)
    resultado = q.aplicar_estrategia(estado_inicial, condiciones, alcance_bin, mecanismo_bin)
    return str(resultado.particion), float(resultado.perdida), float(resultado.tiempo_ejecucion)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def leer_metadatos_hoja(ws) -> tuple[str, str]:
    estado_inicial = str(ws.cell(1, 2).value).strip()
    sistema        = str(ws.cell(2, 2).value).strip()
    return estado_inicial, sistema


def leer_casos_prueba(ws) -> list[dict]:
    casos = []
    for row_idx in range(HEADER_ROW + 1, ws.max_row + 1):
        alc = ws.cell(row_idx, 2).value
        mec = ws.cell(row_idx, 3).value
        if alc is None or mec is None:
            continue
        alc_s = str(alc).strip()
        mec_s = str(mec).strip()
        if not alc_s or not mec_s:
            continue
        casos.append({"row_idx": row_idx, "alcance_letras": alc_s, "mecanismo_letras": mec_s})
    return casos


def letras_a_binario(letras: str, sistema: str) -> str:
    return "".join("1" if c in letras else "0" for c in sistema)


def formatear_tiempo(tiempo_s: float) -> str:
    return (
        f"Horas: {tiempo_s / 3600:.2f} = "
        f"Minutos: {tiempo_s / 60:.1f} = "
        f"Segundos: {tiempo_s:.4f}"
    )


def extraer_pagina(sheet_name: str) -> str:
    m = re.match(r"\d+([A-Z])", sheet_name)
    if not m:
        raise ValueError(f"No se pudo extraer la página del nombre de hoja: {sheet_name!r}")
    return m.group(1)


def escribir_resultado(
    excel_path: Path,
    sheet_name: str,
    row_idx: int,
    col_inicio: int,
    particion: str,
    perdida: float,
    tiempo: str,
) -> None:
    wb = openpyxl.load_workbook(excel_path)
    ws = wb[sheet_name]
    ws.cell(row_idx, col_inicio    ).value = particion
    ws.cell(row_idx, col_inicio + 1).value = str(round(float(perdida), 4))
    ws.cell(row_idx, col_inicio + 2).value = tiempo
    wb.save(excel_path)
    wb.close()


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------
def main() -> None:
    print(f"Excel      : {EXCEL_PATH}")
    print(f"Hoja       : {SHEET_NAME}")
    print(f"k          : 2  (bipartición QNodes)")
    print(f"Workers    : {MAX_WORKERS}  (procesos paralelos)")
    print(f"Cols       : {COL_INICIO}–{COL_INICIO + 2} (Partición, Pérdida, Tiempo)")
    print()

    # 1 — Leer metadatos y casos del Excel
    wb_r = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_r = wb_r[SHEET_NAME]
    estado_inicial, sistema = leer_metadatos_hoja(ws_r)
    condiciones = "1" * len(estado_inicial)
    casos       = leer_casos_prueba(ws_r)
    filas_rellenas = {
        row_idx
        for row_idx in range(HEADER_ROW + 1, ws_r.max_row + 1)
        if ws_r.cell(row_idx, COL_INICIO).value not in (None, "")
    } if SKIP_SI_RELLENO else set()
    wb_r.close()

    pagina = extraer_pagina(SHEET_NAME)

    print(f"Estado inicial : {estado_inicial}  (N={len(estado_inicial)})")
    print(f"Sistema        : {sistema}")
    print(f"Condiciones    : {condiciones}  (todo 1s)")
    print(f"Casos de prueba: {len(casos)}  (ya rellenos: {len(filas_rellenas)})")
    print(f"TPM            : N{len(estado_inicial)}{pagina}.csv")
    print()
    print("-" * 70)

    errores = 0

    # 2 — Preparar trabajos pendientes
    trabajos: list[dict] = []
    for i, caso in enumerate(casos):
        if caso["row_idx"] in filas_rellenas:
            print(f"[{i + 1}/{len(casos)}] fila={caso['row_idx']} → ya relleno, omitiendo")
            continue
        trabajos.append({
            "idx_global": i,
            "caso": caso,
            "alcance_bin":   letras_a_binario(caso["alcance_letras"],   sistema),
            "mecanismo_bin": letras_a_binario(caso["mecanismo_letras"], sistema),
        })

    if not trabajos:
        print("Nada que procesar — todos los casos ya estaban rellenos.")
        return

    print(f"\nProcesando {len(trabajos)} caso(s) con {MAX_WORKERS} worker(s)…\n")

    # 3 — Ejecutar en paralelo; escribir al Excel en el proceso principal (serial)
    with concurrent.futures.ProcessPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_trabajo = {
            executor.submit(
                _ejecutar_caso_qnodes,
                str(QNODES_DIR),
                str(_qnodes_sp),
                estado_inicial,
                pagina,
                condiciones,
                t["alcance_bin"],
                t["mecanismo_bin"],
            ): t
            for t in trabajos
        }

        completados = 0
        for future in concurrent.futures.as_completed(future_to_trabajo):
            t    = future_to_trabajo[future]
            caso = t["caso"]
            completados += 1

            print(
                f"[{completados}/{len(trabajos)}] fila={caso['row_idx']}  "
                f"Alcance={caso['alcance_letras']}  "
                f"Mecanismo={caso['mecanismo_letras']}"
            )

            try:
                particion_str, perdida, tiempo_ejec = future.result()
                tiempo_str = formatear_tiempo(tiempo_ejec)

                print(f"  Partición : {particion_str}")
                print(f"  Pérdida   : {perdida:.4f}")
                print(f"  {tiempo_str}")

                escribir_resultado(
                    EXCEL_PATH, SHEET_NAME, caso["row_idx"],
                    COL_INICIO, particion_str, perdida, tiempo_str,
                )
                print(f"  ✓ Escrito en Excel\n")

            except Exception as e:
                errores += 1
                sep = "=" * 70
                print(f"\n{sep}")
                print(f"ERROR en fila {caso['row_idx']}")
                print(f"  Alcance  : {caso['alcance_letras']} → {t['alcance_bin']}")
                print(f"  Mecanismo: {caso['mecanismo_letras']} → {t['mecanismo_bin']}")
                print(f"  Excepción: {type(e).__name__}: {e}")
                traceback.print_exc()
                print(sep)
                print("Caso omitido — continuando con los demás.\n")

    print("=" * 70)
    print(f"Completado. Hoja='{SHEET_NAME}'  Procesados={len(trabajos)}  Errores={errores}")


if __name__ == "__main__":
    main()
