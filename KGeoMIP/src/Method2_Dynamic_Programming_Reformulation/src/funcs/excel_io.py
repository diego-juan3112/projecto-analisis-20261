"""Lectura y escritura robusta del Excel de pruebas (corrige el hallazgo N11).

En lugar de índices fijos de hoja y rangos mágicos, todo se localiza por
contenido: la hoja por el prefijo de tamaño/página (tolerando espacios y
acentos), la fila de encabezados por la celda '#Prueba', los bloques por k
por los rótulos 'BIPARTICIONES' / 'N-PARTICIONES', y dentro de cada bloque
la sub-columna de la estrategia ('Geometric' o 'QNodes') por su rótulo.
"""

import re
import unicodedata
from dataclasses import dataclass

from openpyxl.styles import Alignment
from openpyxl.worksheet.worksheet import Worksheet

_RE_K_PARTICION = re.compile(r"(\d+)\s*-\s*PARTICION")
_MAX_FILAS_BUSQUEDA = 30


def _normalizar(texto) -> str:
    """Mayúsculas, sin acentos y sin espacios sobrantes, para comparar rótulos."""
    plano = unicodedata.normalize("NFKD", str(texto))
    plano = "".join(c for c in plano if not unicodedata.combining(c))
    return plano.strip().upper()


def localizar_hoja(libro, red: str) -> Worksheet:
    """Hoja cuyo nombre comienza con el identificador de red (ej. '10A')."""
    objetivo = _normalizar(red)
    for nombre in libro.sheetnames:
        if _normalizar(nombre).startswith(objetivo):
            return libro[nombre]
    raise KeyError(f"No hay hoja para la red '{red}' en {libro.sheetnames}")


@dataclass
class ContextoHoja:
    estado_inicial: str
    sistema: str
    condiciones: str
    fila_primer_caso: int
    columnas_geometric: dict[int, int]  # k -> columna 'Partición' (1-based)


def _fila_encabezados(ws: Worksheet) -> int:
    for fila in range(1, _MAX_FILAS_BUSQUEDA + 1):
        valor = ws.cell(fila, 1).value
        if valor and "#PRUEBA" in _normalizar(valor):
            return fila
    raise ValueError(f"No se encontró la fila '#Prueba' en la hoja {ws.title}")


def _valor_rotulado(ws: Worksheet, rotulo: str) -> str:
    objetivo = _normalizar(rotulo)
    for fila in range(1, _MAX_FILAS_BUSQUEDA + 1):
        valor = ws.cell(fila, 1).value
        if valor and objetivo in _normalizar(valor):
            return str(ws.cell(fila, 2).value).strip()
    raise ValueError(f"No se encontró el rótulo '{rotulo}' en la hoja {ws.title}")


def leer_contexto(ws: Worksheet, estrategia: str = "GEOMETRIC") -> ContextoHoja:
    fila_hdr = _fila_encabezados(ws)
    fila_estrategias = fila_hdr - 1
    fila_bloques = fila_hdr - 2

    bloques: dict[int, int] = {}
    for col in range(1, ws.max_column + 1):
        valor = ws.cell(fila_bloques, col).value
        if not valor:
            continue
        texto = _normalizar(valor)
        coincidencia = _RE_K_PARTICION.search(texto)
        if coincidencia:
            bloques[int(coincidencia.group(1))] = col
        elif "BIPARTICION" in texto:
            bloques[2] = col

    ordenados = sorted(bloques.items(), key=lambda kv: kv[1])
    columnas: dict[int, int] = {}
    for i, (k, col_inicio) in enumerate(ordenados):
        col_fin = ordenados[i + 1][1] - 1 if i + 1 < len(ordenados) else ws.max_column
        for col in range(col_inicio, col_fin + 1):
            valor = ws.cell(fila_estrategias, col).value
            if valor and estrategia in _normalizar(valor):
                columnas[k] = col
                break

    estado = _valor_rotulado(ws, "Estado inicial")
    sistema = _valor_rotulado(ws, "Sistema:")
    candidato = _valor_rotulado(ws, "Sistema Candidato")
    condiciones = "".join("1" if letra in candidato else "0" for letra in sistema)

    return ContextoHoja(
        estado_inicial=estado,
        sistema=sistema,
        condiciones=condiciones,
        fila_primer_caso=fila_hdr + 1,
        columnas_geometric=columnas,
    )


def leer_casos(ws: Worksheet, contexto: ContextoHoja):
    """Itera (fila, alcance_letras, mecanismo_letras) hasta la primera fila vacía."""
    fila = contexto.fila_primer_caso
    while True:
        alcance = ws.cell(fila, 2).value
        if alcance is None or not str(alcance).strip():
            break
        mecanismo = ws.cell(fila, 3).value or ""
        yield fila, str(alcance).strip(), str(mecanismo).strip()
        fila += 1


def letras_a_binario(letras: str, sistema: str) -> str:
    return "".join("1" if letra in letras else "0" for letra in sistema)


def escribir_resultado(
    ws: Worksheet,
    fila: int,
    col_particion: int,
    particion: str,
    perdida: float,
    tiempo_s: float,
) -> None:
    """Escribe la tripleta (Partición, Pérdida, Tiempo) con el formato de la hoja."""
    celda = ws.cell(fila, col_particion)
    celda.value = particion
    celda.alignment = Alignment(wrap_text=True, vertical="center")
    ws.cell(fila, col_particion + 1).value = f"{perdida:.4f}"
    ws.cell(fila, col_particion + 2).value = (
        f"Horas: {tiempo_s/3600:.2f} = Minutos: {tiempo_s/60:.1f} "
        f"= Segundos: {tiempo_s:.4f}"
    )


def celda_ocupada(ws: Worksheet, fila: int, col_particion: int) -> bool:
    valor = ws.cell(fila, col_particion).value
    return valor is not None and str(valor).strip() != ""
